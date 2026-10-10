"""Validação e extração de texto de PDF, DOCX, TXT, Markdown e imagens (OCR opcional)."""
from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger("jarmes.docs")

ALLOWED_EXT = {".pdf", ".docx", ".txt", ".md", ".markdown", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
MAX_UNCOMPRESSED_DOCX = 200 * 1024 * 1024


class ImportRejected(ValueError):
    """Arquivo recusado na validação (tipo, tamanho ou conteúdo suspeito)."""


@dataclass
class Extraction:
    pages: list[tuple[int, str]] = field(default_factory=list)   # (nº da página/trecho, texto)
    warnings: list[str] = field(default_factory=list)
    status: str = "full"                                         # full|partial|unreadable
    unit: str = "página"                                         # página|trecho

    @property
    def chars(self) -> int:
        return sum(len(t) for _, t in self.pages)


def ocr_available() -> bool:
    try:
        import pytesseract  # type: ignore

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def validate_file(path: Path, max_mb: int = 100) -> str:
    """Valida extensão, tamanho e assinatura ('magic bytes'). Devolve a extensão normalizada."""
    path = Path(path)
    if not path.is_file():
        raise ImportRejected("Arquivo não encontrado.")
    ext = path.suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ImportRejected(f"Formato '{ext or 'sem extensão'}' não suportado. Use PDF, DOCX, TXT, MD ou imagem.")
    size = path.stat().st_size
    if size == 0:
        raise ImportRejected("O arquivo está vazio.")
    if size > max_mb * 1024 * 1024:
        raise ImportRejected(f"O arquivo tem {size / 1024 / 1024:.0f} MB e excede o limite de {max_mb} MB.")
    head = path.open("rb").read(16)
    if ext == ".pdf" and not head.startswith(b"%PDF"):
        raise ImportRejected("O conteúdo não corresponde a um PDF válido.")
    if ext == ".docx":
        if not head.startswith(b"PK"):
            raise ImportRejected("O conteúdo não corresponde a um DOCX válido.")
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
                if "word/document.xml" not in names:
                    raise ImportRejected("DOCX inválido: documento principal ausente.")
                if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED_DOCX:
                    raise ImportRejected("DOCX recusado: tamanho descompactado suspeito.")
        except zipfile.BadZipFile:
            raise ImportRejected("O DOCX está corrompido.") from None
    if ext in (".png",) and not head.startswith(b"\x89PNG"):
        raise ImportRejected("O conteúdo não corresponde a um PNG válido.")
    if ext in (".jpg", ".jpeg") and not head.startswith(b"\xff\xd8"):
        raise ImportRejected("O conteúdo não corresponde a um JPEG válido.")
    if ext in (".txt", ".md", ".markdown"):
        sample = path.open("rb").read(8192)
        if b"\x00" in sample or sample.startswith((b"MZ", b"\x7fELF", b"PK")):
            raise ImportRejected("O arquivo de texto contém dados binários e foi recusado.")
    return ext


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_units(text: str, size: int = 3000) -> list[tuple[int, str]]:
    paras = re.split(r"\n\s*\n", text)
    units, buf = [], ""
    for p in paras:
        if len(buf) + len(p) > size and buf:
            units.append(buf)
            buf = ""
        buf += (("\n\n" if buf else "") + p)
    if buf.strip():
        units.append(buf)
    return [(i + 1, u.strip()) for i, u in enumerate(units) if u.strip()]


def extract_pdf(path: Path, ocr: bool = True) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    out = Extraction(unit="página")
    try:
        reader = PdfReader(str(path))
    except (PdfReadError, ValueError, OSError) as exc:
        out.status = "unreadable"
        out.warnings.append(f"Não foi possível abrir o PDF: {type(exc).__name__}.")
        return out
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            out.status = "unreadable"
            out.warnings.append("PDF protegido por senha: não foi possível ler o conteúdo.")
            return out
    empty_pages, failed = [], []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = _clean(page.extract_text() or "")
        except Exception:
            text = ""
            failed.append(i)
        if len(text) < 15:
            empty_pages.append(i)
        else:
            out.pages.append((i, text))
    if empty_pages and ocr and ocr_available():
        recovered = _ocr_pdf_pages(path, empty_pages)
        for pno, txt in recovered.items():
            out.pages.append((pno, txt))
        empty_pages = [p for p in empty_pages if p not in recovered]
        out.pages.sort()
        if recovered:
            out.warnings.append(f"OCR aplicado em {len(recovered)} página(s) digitalizada(s); confira o texto.")
    total = len(reader.pages)
    if empty_pages:
        shown = ", ".join(map(str, empty_pages[:15])) + ("…" if len(empty_pages) > 15 else "")
        hint = "" if ocr_available() else " Instale o Tesseract e o pacote pytesseract para habilitar OCR."
        out.warnings.append(f"{len(empty_pages)} de {total} página(s) sem texto extraível (provável digitalização): {shown}.{hint}")
    if failed:
        out.warnings.append(f"Falha ao ler as páginas: {failed[:10]}.")
    if not out.pages:
        out.status = "unreadable"
    elif empty_pages or failed:
        out.status = "partial"
    return out


def _ocr_pdf_pages(path: Path, pages: list[int]) -> dict[int, str]:
    result: dict[int, str] = {}
    try:
        import pypdfium2 as pdfium  # type: ignore
        import pytesseract  # type: ignore

        pdf = pdfium.PdfDocument(str(path))
        for p in pages[:200]:
            image = pdf[p - 1].render(scale=2.0).to_pil()
            text = _clean(pytesseract.image_to_string(image, lang="por"))
            if len(text) >= 15:
                result[p] = text
    except Exception as exc:
        log.warning("OCR de PDF falhou: %s", type(exc).__name__)
    return result


def extract_docx(path: Path) -> Extraction:
    import docx  # python-docx

    out = Extraction(unit="trecho")
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        out.status = "unreadable"
        out.warnings.append(f"Não foi possível abrir o DOCX: {type(exc).__name__}.")
        return out
    parts = []
    for p in document.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.lower().startswith(("heading", "título")):
            t = "# " + t
        parts.append(t)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _clean("\n\n".join(parts))
    if document.inline_shapes and len(document.inline_shapes):
        out.warnings.append(f"O documento contém {len(document.inline_shapes)} imagem(ns) que não foram lidas (sem OCR).")
        out.status = "partial"
    if not text:
        out.status = "unreadable"
        out.warnings.append("Nenhum texto encontrado no DOCX.")
        return out
    out.pages = _split_units(text)
    return out


def extract_text_file(path: Path) -> Extraction:
    raw = path.read_bytes()
    out = Extraction(unit="trecho")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            if enc != "utf-8-sig":
                out.warnings.append(f"Codificação não UTF-8; interpretado como {enc}.")
            break
        except UnicodeDecodeError:
            continue
    text = _clean(text)
    if not text:
        out.status = "unreadable"
        out.warnings.append("Arquivo sem texto.")
        return out
    out.pages = _split_units(text)
    return out


def extract_image(path: Path) -> Extraction:
    out = Extraction(unit="página")
    if not ocr_available():
        out.status = "unreadable"
        out.warnings.append("Imagem armazenada, mas o OCR não está disponível. Instale o Tesseract (idioma 'por') e pytesseract.")
        return out
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore

        text = _clean(pytesseract.image_to_string(Image.open(path), lang="por"))
    except Exception as exc:
        out.status = "unreadable"
        out.warnings.append(f"Falha no OCR: {type(exc).__name__}.")
        return out
    if len(text) < 15:
        out.status = "unreadable"
        out.warnings.append("O OCR não encontrou texto legível na imagem.")
        return out
    out.pages = [(1, text)]
    out.warnings.append("Texto obtido por OCR; confira possíveis erros de reconhecimento.")
    out.status = "partial"
    return out


def extract(path: Path, ext: str | None = None, ocr: bool = True) -> Extraction:
    ext = (ext or path.suffix).lower()
    if ext == ".pdf":
        return extract_pdf(path, ocr=ocr)
    if ext == ".docx":
        return extract_docx(path)
    if ext in IMAGE_EXT:
        return extract_image(path)
    return extract_text_file(path)
