"""Fábricas de arquivos e servidores falsos para os testes."""
from __future__ import annotations

import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path


def make_text_pdf(path: Path, pages: list[str]) -> Path:
    """PDF mínimo e válido com texto extraível (uma linha por página)."""
    objs = []
    n_pages = len(pages)
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(n_pages))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    font_id = 3 + n_pages * 2
    for i, text in enumerate(pages):
        content_id = 4 + i * 2
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def make_image_pdf(path: Path) -> Path:
    from PIL import Image
    Image.new("RGB", (300, 300), "white").save(path, "PDF")
    return path


def make_docx(path: Path, paragraphs: list[str], heading: str | None = None) -> Path:
    import docx
    d = docx.Document()
    if heading:
        d.add_heading(heading, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(path)
    return path


class FakeServer:
    """Servidor HTTP local programável: routes = {(método, caminho): (status, corpo, content-type)}."""

    def __init__(self, routes=None):
        self.routes = routes or {}
        self.requests: list[dict] = []
        outer = self

        class H(BaseHTTPRequestHandler):
            def _do(self):
                length = int(self.headers.get("Content-Length") or 0)
                body = self.rfile.read(length) if length else b""
                path = self.path.split("?")[0]
                outer.requests.append(dict(method=self.command, path=self.path, headers={k.lower(): v for k, v in self.headers.items()}, body=body))
                handler = outer.routes.get((self.command, path))
                if handler is None:
                    status, payload, ctype = 404, b"{}", "application/json"
                elif callable(handler):
                    status, payload, ctype = handler(self.path, body)
                else:
                    status, payload, ctype = handler
                if isinstance(payload, (dict, list)):
                    payload = json.dumps(payload).encode()
                elif isinstance(payload, str):
                    payload = payload.encode()
                self.send_response(status)
                self.send_header("Content-Type", ctype)
                self.send_header("Content-Length", str(len(payload)))
                self.end_headers()
                self.wfile.write(payload)

            do_GET = do_POST = _do

            def log_message(self, *a):
                pass

        self.httpd = HTTPServer(("127.0.0.1", 0), H)
        self.url = f"http://127.0.0.1:{self.httpd.server_port}"
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self, *a):
        self.httpd.shutdown()
        self.httpd.server_close()
